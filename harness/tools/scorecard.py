"""점수표(스코어카드) 생성: 측정 파일만 읽어 QUESTIONS.md Q2 / Q2-a 규칙으로 점수·약점·판정을 계산한다.

    python /repo/harness/tools/scorecard.py            # → reports/scorecard.md, reports/scorecard.xlsx

숫자를 손으로 옮기지 않는다. 입력 파일이 없는 칸은 '미측정'.
"""
import collections
import glob
import json
import os
import re
import statistics as st

REPO = "/repo" if os.path.isdir("/repo/experiments") else "."
EXPW = "/experiments" if os.path.isdir("/experiments") else f"{REPO}/experiments"
# 도구 컨테이너는 저장소를 읽기 전용으로 마운트하므로 쓰기 가능한 experiments 아래에 만들고 호스트가 reports/ 로 복사한다
OUTDIR = os.environ.get("SCORE_OUT", f"{EXPW}/_scorecard")


def jload(p):
    return json.load(open(p, encoding="utf-8"))


# ── L4 ─────────────────────────────────────────────────────────────
L4 = {  # 후보키: (표시 이름, 이미지 크기 MB, 컨테이너 수, 코드·설정 줄 수 산출 파일들, 라이선스)
    "flinksql": ("V1 Flink 1.20.1 SQL", None, 2, ["flink/sql/01_sources.sql", "flink/sql/02_tier1_rules.sql",
                                                  "flink/sql/03_tier1_zscore.sql", "flink/sql/04_tier1_cep.sql"], "Apache-2.0"),
    "flink22": ("Flink 2.2.1 SQL", None, 2, ["flink/sql/01_sources.sql", "flink/sql/02_tier1_rules.sql",
                                             "flink/sql/03_tier1_zscore.sql", "flink/sql/04_tier1_cep.sql",
                                             "candidates/l4-flink22/Dockerfile"], "Apache-2.0"),
    "cep": ("Flink 2.2.1 DataStream CEP", None, 2, ["flink/sql/01_sources.sql", "flink/sql/02_tier1_rules.sql",
                                                     "flink/sql/03_tier1_zscore.sql",
                                                     "candidates/l4-flink-cep/src/main/java/exp/CepJob.java",
                                                     "candidates/l4-flink-cep/pom.xml", "candidates/l4-flink22/Dockerfile"],
            "Apache-2.0"),
    "python": ("Python 서비스", None, 1, ["candidates/l4-python/app.py", "candidates/l4-python/Dockerfile"],
               "PSF-2.0 · Apache-2.0(confluent-kafka)"),
}
IMAGE_MB = json.loads(os.environ.get("IMAGE_MB", "{}"))   # {"flinksql": 1860, ...} 실행 스크립트가 docker images 로 채움


def loc(files):
    n = 0
    for f in files:
        p = f"{REPO}/{f}"
        if os.path.exists(p):
            n += sum(1 for line in open(p, encoding="utf-8") if line.strip())
    return n


def l4_normal(runs):
    acc = collections.defaultdict(lambda: {"pass": 0, "fail": 0, "diff": 0, "p95": [], "mem": [], "cpu": []})
    for run in runs:
        s = jload(f"{EXPW}/EXP-L4/raw/summary_{run}.json")
        base = next(k for k in s if k.startswith("diff_vs_")).replace("diff_vs_", "")
        for c in s["alerts"]:
            r = jload(f"{EXPW}/EXP-L4/raw/result_{run}_{c}.json")
            a = acc[c]
            a["pass"] += r["pass"]
            a["fail"] += r["fail"]
            if c != base:
                d = s[f"diff_vs_{base}"][c]
                a["diff"] += d[f"only_{c}"] + d[f"only_{base}"]
            if r["cep_latency_ms"]["p95"] is not None:
                a["p95"].append(r["cep_latency_ms"]["p95"])
            if c in s["resources"]:
                a["mem"].append(s["resources"][c]["mem_avg_mib"])
                a["cpu"].append(s["resources"][c]["cpu_avg_pct"])
    return acc, base


def l4_restart():
    out = collections.defaultdict(lambda: {"runs": 0, "dup": 0, "lost": 0, "kinds": collections.Counter()})
    for f in glob.glob(f"{EXPW}/EXP-S09/raw/g8_*.json"):
        g = jload(f)
        run = g["run"]
        if not re.match(r"^[xyz]\d+$", run):      # 반복 행렬(x·y·z)만 집계. r13~r15 는 1회성 예비 실행
            continue
        acts = open(f"{EXPW}/EXP-S09/raw/s09_{run}_actions.log", encoding="utf-8").read()
        kind = "JM+TM" if "jobmanager" in acts else ("TM" if "taskmanager" in acts else "process")
        o = out[g["killed"]]
        o["runs"] += 1
        o["dup"] += g["duplicates"]
        o["lost"] += g["lost_vs_control"]
        o["kinds"][kind] += 1
    return out


def score(vals, lower_better=True):
    """vals: {cand: value or None} → {cand: score 1~5 or None}"""
    xs = {k: v for k, v in vals.items() if v is not None}
    if not xs:
        return {k: None for k in vals}
    best = min(xs.values()) if lower_better else max(xs.values())
    res = {}
    for k, v in vals.items():
        if v is None:
            res[k] = None
        elif lower_better:
            res[k] = 5.0 if v == best else max(1.0, round(5 * (best / v), 1)) if v else 5.0
        else:
            res[k] = 5.0 if v == best else max(1.0, round(5 * (v / best), 1)) if best else 5.0
    return res


def build_l4():
    runs = sorted({re.search(r"summary_(m\d+)\.json", f).group(1) for f in glob.glob(f"{EXPW}/EXP-L4/raw/summary_m*.json")})
    four = [r for r in runs if set(jload(f"{EXPW}/EXP-L4/raw/summary_{r}.json")["alerts"]) >= set(L4)]
    use = four or runs
    acc, base = l4_normal(use)
    rst = l4_restart()
    rows = {}
    for c, (name, _, cont, files, lic) in L4.items():
        a = acc.get(c)
        r = rst.get(c)
        rows[c] = {
            "name": name,
            "correct": f"{a['pass']}/{a['pass'] + a['fail']}" if a else None,
            "fail": a["fail"] if a else None,
            "diff": a["diff"] if a else None,
            "p95": round(st.mean(a["p95"])) if a and a["p95"] else None,
            "recovery": (r["dup"] + r["lost"]) if r else None,
            "rec_detail": (f"{r['runs']}회({dict(r['kinds'])}) 중복 {r['dup']}·유실 {r['lost']}") if r else None,
            "cpu": round(st.mean(a["cpu"]), 1) if a and a["cpu"] else None,
            "mem": round(st.mean(a["mem"])) if a and a["mem"] else None,
            "image": IMAGE_MB.get(c),
            "containers": cont,
            "loc": loc(files),
            "license": lic,
        }
    sc = {
        "speed": score({c: v["p95"] for c, v in rows.items()}),
        "cpu": score({c: v["cpu"] for c, v in rows.items()}),
        "mem": score({c: v["mem"] for c, v in rows.items()}),
        "image": score({c: v["image"] for c, v in rows.items()}),
        "containers": score({c: v["containers"] for c, v in rows.items()}),
        "loc": score({c: v["loc"] for c, v in rows.items()}),
    }
    b = rows["flinksql"]
    for c, v in rows.items():
        # 정확도·복구: 기준선보다 나쁘면 0 (Q2-a)
        worse_acc = v["fail"] is None or (b["fail"] is not None and v["fail"] > b["fail"]) or (v["diff"] or 0) > 0
        worse_rec = v["recovery"] is None or (b["recovery"] is not None and v["recovery"] > b["recovery"])
        v["s_acc"] = None if v["fail"] is None else (0 if worse_acc else 5)
        best_rec = min(x["recovery"] for x in rows.values() if x["recovery"] is not None) if any(
            x["recovery"] is not None for x in rows.values()) else None
        v["s_rec"] = None if v["recovery"] is None else (0 if worse_rec else (5.0 if v["recovery"] == best_rec else
                                                                             max(1.0, round(5 * ((best_rec + 1) / (v["recovery"] + 1)), 1))))
        for k in sc:
            v[f"s_{k}"] = sc[k][c]
        parts = [v[k] for k in ("s_acc", "s_speed", "s_rec", "s_cpu", "s_mem", "s_image", "s_containers", "s_loc")]
        v["total"] = round(sum(p for p in parts if p is not None), 1) if all(p is not None for p in parts) else None
        v["missing"] = [k for k, p in zip(("정확도", "속도", "복구", "CPU", "메모리", "이미지", "컨테이너", "줄 수"), parts) if p is None]
        # 판정 (Q2)
        if v["missing"]:
            v["verdict"] = "미완료(미측정: " + ", ".join(v["missing"]) + ")"
        elif c == "flinksql":
            v["verdict"] = "기준선"
        elif worse_acc or worse_rec or (v["p95"] > b["p95"]):
            bad = [n for n, w in (("정확도", worse_acc), ("복구", worse_rec), ("속도", v["p95"] > b["p95"])) if w]
            v["verdict"] = "탈락(기준선보다 나쁨: " + ", ".join(bad) + ")"
        else:
            v["verdict"] = "후보(성능 ≥ 기준선)"
        # 약점: 각 지표에서 최선 대비
        weak = []
        best = {k: min(x[k] for x in rows.values() if x[k] is not None) for k in ("p95", "cpu", "mem", "image", "loc")
                if any(x[k] is not None for x in rows.values())}
        if v["recovery"]:
            weak.append(f"재시작 {v['rec_detail']}")
        for k, lab, unit in (("p95", "지연 p95", "ms"), ("cpu", "CPU", "%"), ("mem", "메모리", "MiB"), ("image", "이미지", "MB"), ("loc", "줄 수", "")):
            if v[k] is not None and k in best and best[k] and v[k] >= best[k] * 1.5:
                weak.append(f"{lab} {v[k]}{unit} (최선 {best[k]}{unit}의 {v[k] / best[k]:.0f}배)")
        v["weak"] = " · ".join(weak) if weak else "—"
    # 성능 동점 후보 중 선택 (Q2-3·4)
    ok = [c for c, v in rows.items() if v["verdict"].startswith("후보")]
    if ok:
        best_perf = sorted(ok, key=lambda c: (rows[c]["recovery"], rows[c]["p95"]))
        pick = best_perf[0]
        for c in ok:
            rows[c]["verdict"] = "채택" if c == pick else "후보(성능 차순위)"
    return rows, use


def fmt(v, suf=""):
    return "미측정" if v is None else f"{v}{suf}"


def md_l4(rows, runs):
    h = ("| 후보 | 정확도 | 속도 p95 | 복구(유실+중복) | CPU 평균 | 메모리 평균 | 이미지 | 컨테이너 | 코드·설정 줄 | 합계(40) | 판정 | 약점 |\n"
         "|---|---|---|---|---|---|---|---|---|---|---|---|\n")
    out = []
    for c, v in rows.items():
        cell = lambda val, s, suf="": f"{fmt(val, suf)} **{'-' if s is None else s}**"
        out.append(f"| {v['name']} | {fmt(v['correct'])} · V1 대비 차이 {fmt(v['diff'])} **{'-' if v['s_acc'] is None else v['s_acc']}** "
                   f"| {cell(v['p95'], v['s_speed'], ' ms')} | {cell(v['recovery'], v['s_rec'])} | {cell(v['cpu'], v['s_cpu'], '%')} "
                   f"| {cell(v['mem'], v['s_mem'], ' MiB')} | {cell(v['image'], v['s_image'], ' MB')} | {cell(v['containers'], v['s_containers'])} "
                   f"| {cell(v['loc'], v['s_loc'])} | {fmt(v['total'])} | {v['verdict']} | {v['weak']} |")
    return (f"측정 실행: {', '.join(runs)} (정상), EXP-S09 x·y·z 행렬(재시작). 굵은 숫자 = 5점 만점 점수.\n\n" + h + "\n".join(out) + "\n")


def xlsx(sheets, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    wb.remove(wb.active)
    for title, header, data in sheets:
        ws = wb.create_sheet(title)
        ws.append(header)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="305496")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for row in data:
            ws.append(row)
        for col in ws.columns:
            width = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(60, max(10, width + 2))
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if isinstance(cell.value, str) and cell.value.startswith("탈락"):
                    cell.fill = PatternFill("solid", fgColor="F8CBAD")
                elif cell.value == "채택":
                    cell.fill = PatternFill("solid", fgColor="C6EFCE")
        ws.freeze_panes = "B2"
    wb.save(path)


def main():
    rows, runs = build_l4()
    md = "# 점수표 (자동 생성 — harness/tools/scorecard.py)\n\n규칙: QUESTIONS.md Q2·Q2-a. 합계는 읽기 편의용, 판정은 Q2 규칙.\n\n## L4 이상탐지\n\n" + md_l4(rows, runs)
    os.makedirs(OUTDIR, exist_ok=True)
    open(f"{OUTDIR}/scorecard.md", "w", encoding="utf-8").write(md)
    header = ["후보", "정확도", "V1 대비 알람 차이", "정확도 점수", "지연 p95(ms)", "속도 점수", "복구(유실+중복)", "복구 상세", "복구 점수",
              "CPU 평균(%)", "CPU 점수", "메모리 평균(MiB)", "메모리 점수", "이미지(MB)", "이미지 점수", "컨테이너", "컨테이너 점수",
              "코드·설정 줄", "줄 수 점수", "합계(40)", "판정", "약점", "라이선스"]
    data = [[v["name"], fmt(v["correct"]), fmt(v["diff"]), v["s_acc"], v["p95"], v["s_speed"], v["recovery"], v["rec_detail"], v["s_rec"],
             v["cpu"], v["s_cpu"], v["mem"], v["s_mem"], v["image"], v["s_image"], v["containers"], v["s_containers"],
             v["loc"], v["s_loc"], v["total"], v["verdict"], v["weak"], v["license"]] for v in rows.values()]
    sheets = [("L4 이상탐지", header, data)]
    extra = os.environ.get("EXTRA_SHEETS")
    if extra and os.path.exists(extra):
        for s in jload(extra):
            sheets.append((s["title"], s["header"], s["rows"]))
            md += f"\n## {s['title']}\n\n" + "| " + " | ".join(s["header"]) + " |\n|" + "---|" * len(s["header"]) + "\n" + \
                  "\n".join("| " + " | ".join(fmt(x) for x in r) + " |" for r in s["rows"]) + "\n"
        open(f"{OUTDIR}/scorecard.md", "w", encoding="utf-8").write(md)
    xlsx(sheets, f"{OUTDIR}/scorecard.xlsx")
    print(md)


if __name__ == "__main__":
    main()
