"""Build the application plus actual-screen attachment without changing its source draft."""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'submission'
OUT.mkdir(exist_ok=True)
doc = Document(ROOT / 'panel-assets/CECO-panel-application-draft.docx')

replacements = {
    '실제 화면 동선은 리허설 후 확정한다.': '전시 동선은 최종 리허설에서 점검한다.',
    '화면 리허설과 최종 캡처는 준비 중이다.': '실제 화면에서 AI 대응안 생성·승인·정지 확인·반려와 새로고침 후 기록 보존, 기존 SCADA 재가동을 확인했다.',
    '시스템 구성도와 실제 데모 화면을 사용한다. 운영 현황, 설비–센서–문서 그래프, 근거·검토, 조치 결과 화면은 실제 캡처 후 첨부한다.': '시스템 구성도와 실제 실행 화면 6장을 제공한다. 지식 검토·게시, 설비 관계, AI 검토·정지 결과, 반려 기록, 기존 SCADA 재가동을 보여준다.',
}
for table in doc.tables:
    table.autofit = False
    table.columns[0].width, table.columns[1].width = Cm(2.5), Cm(15.3)
    width = table._tbl.tblPr.find(qn('w:tblW'))
    width.set(qn('w:type'), 'dxa')
    width.set(qn('w:w'), str(round(17.8 / 2.54 * 1440)))
    for row in table.rows:
        row.cells[0].width, row.cells[1].width = Cm(2.5), Cm(15.3)
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    for before, after in replacements.items():
                        if before in run.text:
                            run.text = run.text.replace(before, after)
for sec in doc.sections:
    for paragraph in sec.footer.paragraphs:
        for run in paragraph.runs:
            run.text = run.text.replace('검토용 초안  |', '패널 신청자료  |')

screens = [
    ('ui-knowledge-reviewed.png', '01  원본 자료에서 관계 후보 검토', '문서와 공정 설정을 실제 업로드·선택해 루나가 의미 관계를 제안했다. 원문 인용과 설비 식별자를 검토하고 사람의 확인 후 게시했다.'),
    ('final-graph.png', '02  게시된 설비·센서·문서 관계', '실제 Neo4j의 36개 노드·45개 관계다. 설비 2개, 센서 12개, 문서 3개와 절·제어점·알람 규칙을 연결한다. 교육용 관계이며 실제 고장 원인 확정이 아니다.'),
    ('ui-review-layout.png', '03  근거 기반 대응안과 승인 후 결과', '루나는 최신 GOOD 전류·진동의 동시 상한 초과와 문서 3개를 조회해 정지를 제안했다. 그림은 승인 완료 후 저장된 분석문과 정지 확인 결과를 다시 연 화면이다.'),
    ('ui-stop-persisted.png', '04  승인·정지 확인·후속 점검 기록', '같은 사건에서 사람 승인 뒤 시뮬레이터 정지를 확인했다. 새로고침 후에도 승인 사유와 처리 이력이 유지된다. 정지 확인은 원인 제거·정비 완료를 뜻하지 않는다.'),
    ('ai-rejection-persisted.png', '05  별도 사건의 반려 경로', '관측 최신성 재확인과 점검 범위 구체화를 사유로 반려한 별도 사건이다. 사유와 이력을 보존하며 이 분기에서는 정지 명령을 실행하지 않았다.'),
    ('scada-after-restart.png', '06  다음 체험을 위한 강사 재가동', '고장 주입 해제 후 기존 FUXA 화면의 M-101 기동 버튼으로 운전을 복구했다. 새 운전 상태와 GOOD 관측을 확인했다. AI에는 재가동 권한을 추가하지 않았다.'),
]
for index, (source, caption, description) in enumerate(screens):
    if index % 2 == 0:
        doc.add_page_break()
        title = ['지식 구축과 관계 탐색', 'AI 검토와 승인 결과', '반려와 다음 체험 준비'][index // 2]
        p = doc.add_paragraph(f'별첨 {index // 2 + 1}  {title}', 'Title')
        p.paragraph_format.space_after = Pt(6)
        p = doc.add_paragraph('실제 실행 화면 · 가상 AR-100 공정 · 2026.09.21 · 각 설명의 실행 단계 기준 · 검토·결과 화면은 관련 영역 확대')
        p.paragraph_format.space_after = Pt(8)
        p.runs[0].font.size = Pt(9)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    crops = {2: (730, 210, 820, 540), 3: (730, 390, 820, 570), 4: (730, 450, 820, 540)}
    if index in crops:
        x,y,w,h = crops[index]
        picture = run.add_picture(str(ROOT / 'browser-qa' / source), width=Cm(14), height=Cm(14*h/w))
        crop = OxmlElement('a:srcRect')
        for key,value in {'l':x/1600,'t':y/1000,'r':(1600-x-w)/1600,'b':(1000-y-h)/1000}.items():
            crop.set(key,str(round(value*100000)))
        picture._inline.xpath('.//pic:blipFill')[0].insert(1,crop)
    else:
        run.add_picture(str(ROOT / 'browser-qa' / source), width=Cm(15.5))
    run._r.xpath('.//wp:docPr')[0].set('descr', description)
    p = doc.add_paragraph(caption)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(10)
    p = doc.add_paragraph(description)
    p.paragraph_format.space_after = Pt(7)
    p.runs[0].font.size = Pt(9)
for p in doc.paragraphs:
    snap = OxmlElement('w:snapToGrid')
    snap.set(qn('w:val'), '0')
    p._p.get_or_add_pPr().append(snap)
doc.core_properties.title = '온톨로지 기반 제조 AI 업무도우미 패널 신청자료'
doc.core_properties.subject = '신청서와 실제 시연 화면 별첨'
doc.core_properties.comments = ''
target = OUT / 'CECO_패널신청자료_시연검증본.docx'
doc.save(target)
print(target)
