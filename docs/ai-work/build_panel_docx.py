"""Editable CECO application draft, based on the supplied two-column example.

Run with the Codex bundled Python. No fabricated service screenshots.
"""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'panel-assets' / 'CECO-panel-application-draft.docx'
doc = Document()
sec = doc.sections[0]
# A4 portrait follows the proportions of the supplied application image.
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(1.7)
sec.left_margin = sec.right_margin = Cm(1.6)
sec.footer_distance = Cm(.7)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    style = doc.styles[name]
    style.font.name = '맑은 고딕'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    for attr in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
        style.element.rPr.rFonts.attrib.pop(qn('w:'+attr), None)
    for border in style.element.xpath('./w:pPr/w:pBdr'):
        border.getparent().remove(border)
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing = Pt(14.5)
doc.styles['Normal'].paragraph_format.space_after = Pt(4)
doc.styles['Title'].font.size = Pt(16)
doc.styles['Title'].font.bold = True
doc.styles['Title'].paragraph_format.space_after = Pt(9)

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = footer.add_run('검토용 초안  |  2026.09.21  |  ')
r.font.size = Pt(8)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

def fill(cell, color):
    x = OxmlElement('w:shd'); x.set(qn('w:fill'), color)
    cell._tc.get_or_add_tcPr().append(x)

def para(cell, text, bold=False, before=0):
    p = cell.paragraphs[0] if len(cell.paragraphs) == 1 and not cell.paragraphs[0].text else cell.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    r = p.add_run(text); r.bold = bold
    snap = OxmlElement('w:snapToGrid'); snap.set(qn('w:val'), '0'); p._p.get_or_add_pPr().append(snap)
    if bold: p.paragraph_format.keep_with_next = True
    return p

def table():
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    t.columns[0].width, t.columns[1].width = Cm(2.5), Cm(15.3)
    pr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        b = OxmlElement('w:'+side)
        for k,v in [('val','single'),('sz','5'),('color','7A8792')]: b.set(qn('w:'+k),v)
        borders.append(b)
    pr.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for side,value in [('top','115'),('bottom','115'),('left','155'),('right','155')]:
        x = OxmlElement('w:'+side); x.set(qn('w:w'),value); x.set(qn('w:type'),'dxa'); margins.append(x)
    pr.append(margins)
    for c,text in zip(t.rows[0].cells,['구분','내용']):
        fill(c,'E6EEF5'); p=para(c,text,True); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    header = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(header)
    return t

def row(t, label):
    rr=t.add_row()
    a,b=rr.cells
    a.width,b.width=Cm(2.5),Cm(15.3)
    fill(a,'DFEAF4'); a.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p=para(a,label,True); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next=False
    return b

doc.add_paragraph('패널 신청내용', 'Title')
t=table()
c=row(t,'비전 및\n목표')
para(c,'작품명  온톨로지 기반 제조 AI 업무도우미',True)
para(c,'교반기 이상 알람을 받으면 센서 이력·설비 정보·대응 문서를 연결해 보여준다. AI가 관측 사실과 미확인 원인, 점검 항목을 정리하고 사용자가 대응안을 검토한다. 승인된 시뮬레이터 조치는 새 관측으로 결과를 확인한다.')
para(c,'비전',True,5)
para(c,'알람 이후 여러 화면과 문서를 찾는 부담을 줄이고, 작업자가 근거를 확인하며 대응을 결정하는 업무 환경을 제공한다.')
para(c,'목표',True,5)
para(c,'설비·센서·문서의 관계를 온톨로지와 그래프 DB로 연결하고, AI 근거 조회부터 승인·반려와 조치 결과 확인까지 통합한다. 가상 반응기 AR-100을 이용한 교육·전시용 강사 시연본이다.')
c=row(t,'성과 등\n주요내용')
para(c,'작동 흐름',True)
para(c,'가상 이상 발생 → SCADA 알람·사건 생성 → 센서 이력·설비 관계·문서 조회 → AI 대응안 → 사용자 승인 또는 반려 → 승인 조건 재확인 → 시뮬레이터 조치 → 결과 확인·후속 점검')
para(c,'주요 기능',True,7)
for text in [
    '반복 알람을 사건으로 묶고, 원본 알람과 발생 시각을 보존한다.',
    '설비 설정·문서에서 관계 후보를 만들고 출처·적용 대상을 검토해 그래프에 게시한다. 게시된 설비–센서–문서 관계를 탐색한다.',
    'AI가 사건 당시 이력과 현재 관측을 구분해 원인 후보와 근거를 정리한다. 문서 부족·충돌, 불량·오래된 관측은 확인할 내용으로 남긴다.',
    '사용자 승인 시 현재 상태와 근거를 다시 확인한다. 반려 사유와 조치 결과를 기록하며, 정지 확인과 정비 완료를 구분한다.',
]: para(c,'• '+text)
para(c,'시스템 구성',True,7)
for text in [
    'IIoT·SCADA: 가상 센서·제어, 데이터 수집, 규칙 기반 이상 감지, 시계열 이력 저장.',
    '온톨로지·Neo4j: 설비·센서·문서 관계, 문서 버전과 출처, 적용 절차 조회.',
    'AI 업무 계층: LiteLLM 중계의 루나 모델이 도구를 조회해 대응안을 작성하고, PostgreSQL에 사건·검토·결과를 보존.',
    '통합 웹: 현황·추세·지식 구축·관계 탐색·검토·결과를 제공. Process GPT와 Ontology Studio의 필요한 코드를 선별·리팩토링해 내부 통합.',
]: para(c,'• '+text)

doc.add_page_break()
t=table()
c=row(t,'성과 등\n주요내용\n계속')
para(c,'관람객 체험',True)
para(c,'강사가 가상 이상을 시작하면 관람객은 사건을 선택해 전류·진동 추세와 대응 문서를 살펴본다. AI 대응안의 근거를 확인하고 승인 또는 반려 흐름을 관찰한다. 정지 후에도 남은 점검 항목을 확인한다. 실제 화면 동선은 리허설 후 확정한다.')
para(c,'성과 확인',True,5)
para(c,'실제 모델·API 실행 기록으로 알람 수신, 지식 조회, 대응안, 승인 후 시뮬레이터 정지와 후속 관측을 확인했다. 문서 부족·충돌, 관측 품질 문제, 모델 연결 실패도 시험했다. 화면 리허설과 최종 캡처는 준비 중이다.')
para(c,'전시 운영',True,5)
para(c,'가상 공정과 교육용 문서임을 안내한다. 강사가 시뮬레이터의 정지·복구를 관리하고 다음 체험 전 고장 주입과 운전 상태를 확인한다. 통신 실패나 결과 불명확 상태에서는 명령을 자동 반복하지 않는다.')
c=row(t,'관련소스')
para(c,'패널 제작용 자료',True)
para(c,'시스템 구성도와 실제 데모 화면을 사용한다. 운영 현황, 설비–센서–문서 그래프, 근거·검토, 조치 결과 화면은 실제 캡처 후 첨부한다.')
p=c.add_paragraph(); p.paragraph_format.space_after=Pt(3)
p.paragraph_format.line_spacing=1.0
snap = OxmlElement('w:snapToGrid'); snap.set(qn('w:val'), '0'); p._p.get_or_add_pPr().append(snap)
r=p.add_run(); r.add_picture(str(ROOT/'panel-assets'/'system-flow.png'),width=Cm(14.5))
inline=r._r.xpath('.//wp:docPr')[0]
inline.set('descr','알람 수집, 게시된 설비와 문서 조회, AI 근거 분석, 사람 승인과 조치 결과 확인의 시스템 구성도. 실제 화면 캡처가 아닌 설명 그림.')
p=para(c,'그림 1  시스템 구성도  ·  실제 서비스 화면 캡처 아님')
p.runs[0].font.size=Pt(8)
para(c,'기술 참고',True,5)
para(c,'INOXPA BCI 매뉴얼 20.005.30.01EN (B), 2024/08, p.15 · IEC 62682:2022 공식 공개 설명 · OPC UA for Machinery 공식 자료. 가상 설비의 운전 한계나 인증 근거로 전용하지 않는다.')

p=doc.add_paragraph('상기와 같이 부스 패널 원고 신청서를 작성하여 제출합니다.')
p.paragraph_format.space_before=Pt(12); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=doc.add_paragraph('2026년        월        일\n담당자                         (인)')
p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
doc.core_properties.title='온톨로지 기반 제조 AI 업무도우미 패널 신청서'
doc.core_properties.subject='CECO 교육 전시 신청서 검토용 초안'
doc.core_properties.author=''
doc.core_properties.comments='실제 화면 캡처 및 UI 리허설 미완료. 검토용 초안.'
for p in doc.paragraphs:
    snap = OxmlElement('w:snapToGrid'); snap.set(qn('w:val'), '0'); p._p.get_or_add_pPr().append(snap)
OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
print(OUT)
